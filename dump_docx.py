# -*- coding: utf-8 -*-
"""Dump 肖如涵简历.docx as readable text, marking bold runs with **."""
import sys

from docx import Document
from docx.oxml.ns import qn

PATH = r"E:\development\AI-Project\肖如涵简历.docx"
doc = Document(PATH)

def run_md(r):
    t = r.text
    if not t:
        return ""
    rPr = r._element.rPr
    rFonts = rPr.find(qn("w:rFonts")) if rPr is not None else None
    ea = (rFonts.get(qn("w:eastAsia")) if rFonts is not None else "") or ""
    bold = "M" in ea
    try:
        sz = r.font.size.pt if r.font.size else None
    except ValueError:
        sz = None
    if bold:
        return f"**{t}**"
    return t

out = []
for i, p in enumerate(doc.paragraphs):
    txt = "".join(run_md(r) for r in p.runs)
    if txt.strip():
        out.append(f"[{i:02d}] {txt}")

for ti, t in enumerate(doc.tables):
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                txt = "".join(run_md(r) for r in p.runs)
                if txt.strip():
                    out.append(f"[T{ti} {ri},{ci}] {txt}")

sys.stdout.reconfigure(encoding="utf-8")
print("\n".join(out))
print(f"\n--- {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables ---")
