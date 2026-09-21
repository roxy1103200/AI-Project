# -*- coding: utf-8 -*-
"""Dump the generated docx's layout parameters so they can be checked against the PDF."""
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

DOCX = r"E:\development\AI-Project\直聘简历-未命名.docx"
doc = Document(DOCX)

print("=== TARGETS measured off the source PDF ===")
print("  page 595.5 x 842.25 pt | margins L/R 24pt, top 23.6pt, bottom 16.35pt")
print("  name 22.5pt #000 | section band 15pt #000 fill F2F3F5 lead 20.25pt")
print("  entry title 12pt | body 9.75pt #333333 lead 18pt exact | date 9.75pt #666666")

sec = doc.sections[0]
print("\n=== SECTION ===")
print(f"  page   : {sec.page_width.pt:.2f} x {sec.page_height.pt:.2f} pt")
print(f"  margins: L={sec.left_margin.pt:.2f} R={sec.right_margin.pt:.2f} "
      f"T={sec.top_margin.pt:.2f} B={sec.bottom_margin.pt:.2f}")
print(f"  content width: {sec.page_width.pt - sec.left_margin.pt - sec.right_margin.pt:.2f} pt")


def runinfo(r):
    rPr = r._element.rPr
    rFonts = rPr.find(qn("w:rFonts")) if rPr is not None else None
    ea = rFonts.get(qn("w:eastAsia")) if rFonts is not None else None
    col = r.font.color.rgb if r.font.color and r.font.color.type is not None else None
    txt = r.text if r.text.strip() else f"<img:{len(r._element.findall(qn('w:drawing')))}>"
    try:                                    # python-docx rejects fractional half-points
        size = f"{r.font.size.pt}pt" if r.font.size else "?"
    except ValueError:
        el = rPr.find(qn("w:sz")) if rPr is not None else None
        size = el.get(qn("w:val")) if el is not None else "?"
    return f"{size} {ea} #{col} | {txt[:34]!r}"


print("\n=== BODY PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    pf = p.paragraph_format
    pPr = p._p.pPr
    shd = pPr.find(qn("w:shd")) if pPr is not None else None
    fill = shd.get(qn("w:fill")) if shd is not None else None
    ls = pf.line_spacing.pt if hasattr(pf.line_spacing, "pt") else pf.line_spacing
    tabs = [f"{t.position.pt:.1f}/{t.alignment}" for t in pf.tab_stops]
    flags = []
    if fill:
        flags.append(f"shd={fill}")
    if ls:
        flags.append(f"lead={ls:.2f}")
    if pf.space_before and pf.space_before.pt:
        flags.append(f"sb={pf.space_before.pt:.1f}")
    if pf.space_after and pf.space_after.pt:
        flags.append(f"sa={pf.space_after.pt:.1f}")
    if pf.first_line_indent and pf.first_line_indent.pt:
        flags.append(f"fli={pf.first_line_indent.pt:.1f}")
    if tabs:
        flags.append(f"tabs={tabs}")
    print(f"\n[{i:02d}] {' '.join(flags)}")
    for r in p.runs:
        print(f"     {runinfo(r)}")

print("\n=== TABLE (header) ===")
for t in doc.tables:
    print(f"  cols: {[c.width.pt for c in t.columns]}")
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            print(f"  cell[{ri}][{ci}] width={cell.width.pt:.2f}")
            for p in cell.paragraphs:
                pf = p.paragraph_format
                ls = pf.line_spacing.pt if hasattr(pf.line_spacing, "pt") else pf.line_spacing
                print(f"     align={p.alignment} lead={ls} sb={pf.space_before}")
                for r in p.runs:
                    print(f"       {runinfo(r)}")
