# -*- coding: utf-8 -*-
"""Rebuild 直聘简历-未命名.pdf as an editable .docx, matching the original layout.

Geometry was measured off the source PDF with PyMuPDF:
  page A4 595.5x842.25pt, content column x=24..571.5 (547.5pt wide)
  section bands 20.25pt tall, fill #F2F3F5
  body 9.75pt / exact 18pt leading / #333333 ; dates #666666
"""
import os

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "直聘简历-未命名.docx")
ASSETS = os.path.join(HERE, "_assets")

FONT_R = "Alibaba PuHuiTi R"
FONT_M = "Alibaba PuHuiTi M"          # heavier weight, stands in for "Bold"
BODY = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x66, 0x66, 0x66)
INK = RGBColor(0x00, 0x00, 0x00)
BAND = "F2F3F5"

CONTENT_W_PT = 547.5
TITLE_COL_PT = 459.8                   # left header cell; centres text at x=253.9
PHOTO_COL_PT = 87.75


# ---------------------------------------------------------------- helpers
def style_run(run, size, color=BODY, bold=False, font=None):
    """Apply font + colour. Both ascii and eastAsia are set so CJK keeps the face."""
    name = font or (FONT_M if bold else FONT_R)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = False              # weight comes from the family, not synthesis
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)

    # w:sz holds integer half-points, so Pt(9.75) silently truncates to 9.5pt.
    # ST_HpsMeasure also admits a universal measure, which is the only
    # schema-valid way to say 9.75pt -- needed because line breaks depend on
    # glyph advance width.
    half_points = size * 2
    if half_points != int(half_points):
        for tag in ("w:sz", "w:szCs"):
            el = rPr.find(qn(tag))
            if el is not None:
                el.set(qn("w:val"), f"{size:g}pt")
    return run


def shade(paragraph, fill=BAND):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def leading(paragraph, pt, before=0, after=0):
    pf = paragraph.paragraph_format
    pf.line_spacing = Pt(pt)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    return paragraph


def set_grid(table, widths_pt):
    """Pin tblGrid + fixed layout.

    Setting cell widths alone leaves the grid equally divided, and Word lays out
    against the grid, so the columns would come out wrong.
    """
    tbl = table._tbl
    old = tbl.find(qn("w:tblGrid"))
    if old is not None:
        tbl.remove(old)
    grid = OxmlElement("w:tblGrid")
    for w in widths_pt:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(round(w * 20))))   # pt -> twips
        grid.append(gc)
    tbl.tblPr.addnext(grid)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl.tblPr.append(layout)


def clear_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        borders.append(el)
    tblPr.append(borders)

    # zero cell padding so the photo lands flush on the right margin
    mar = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), "0")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def body_para(doc, runs, indent=0.0):
    """One 9.75pt/18pt body paragraph. `runs` is a list of (text, bold) tuples."""
    p = doc.add_paragraph()
    leading(p, 18)
    if indent:
        p.paragraph_format.left_indent = Pt(indent)
    for text, bold in runs:
        style_run(p.add_run(text), 9.75, BODY, bold)
    return p


def section_band(doc, title, before=24.6):
    """Full-width grey bar. firstLine indent keeps shading wide while text sits at x=30."""
    p = doc.add_paragraph()
    leading(p, 20.25, before=before, after=7.6)
    p.paragraph_format.first_line_indent = Pt(6)
    shade(p)
    style_run(p.add_run(title), 15, INK, bold=True)
    return p


def entry_title(doc, title, date, before=0.0, after=0.0):
    """12pt bold name with a right-aligned 9.75pt date on the same line."""
    p = doc.add_paragraph()
    leading(p, 20, before=before, after=after)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(CONTENT_W_PT), WD_TAB_ALIGNMENT.RIGHT)
    style_run(p.add_run(title), 12, BODY, bold=True)
    style_run(p.add_run("\t"), 9.75, MUTED)
    style_run(p.add_run(date), 9.75, MUTED)
    return p


# ---------------------------------------------------------------- document
doc = Document()

# normal style reset so nothing inherits Word's defaults
normal = doc.styles["Normal"]
normal.font.name = FONT_R
normal.font.size = Pt(9.75)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_R)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)

sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = Pt(24)
sec.right_margin = Pt(24)
sec.top_margin = Pt(23.6)
sec.bottom_margin = Pt(16.35)

# ---- header: centred name block (left cell) + photo (right cell)
table = doc.add_table(rows=1, cols=2)
table.autofit = False
clear_borders(table)
left, right = table.rows[0].cells
left.width, right.width = Pt(TITLE_COL_PT), Pt(PHOTO_COL_PT)
left.vertical_alignment = WD_ALIGN_VERTICAL.TOP
right.vertical_alignment = WD_ALIGN_VERTICAL.TOP
set_grid(table, (TITLE_COL_PT, PHOTO_COL_PT))

name = left.paragraphs[0]
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
leading(name, 32)
style_run(name.add_run("肖如涵"), 22.5, INK, bold=True)

contact = left.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
leading(contact, 18, before=8.4)
style_run(contact.add_run("男 | 年龄：23岁 | "), 9.75, BODY)
contact.add_run().add_picture(os.path.join(ASSETS, "img0_xref8.png"), width=Pt(12))
style_run(contact.add_run("13177744002 | "), 9.75, BODY)
contact.add_run().add_picture(os.path.join(ASSETS, "img1_xref11.png"), width=Pt(12))
style_run(contact.add_run("roxy1103200@gmail.com"), 9.75, BODY)

intent = left.add_paragraph()
intent.alignment = WD_ALIGN_PARAGRAPH.CENTER
leading(intent, 18)
style_run(intent.add_run("求职意向：agent开发工程师 | 期望城市：广州"), 9.75, BODY)

photo_para = right.paragraphs[0]
photo_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
# single spacing, not "exactly 72.75pt": Word clips an inline image that fills
# an exact-height line. Single spacing lets the line grow to the image.
photo_para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
photo_para.paragraph_format.space_before = Pt(0)
photo_para.paragraph_format.space_after = Pt(0)
photo_para.add_run().add_picture(os.path.join(ASSETS, "img2_xref13.png"), width=Pt(72.75))

# ---- 个人优势
section_band(doc, "个人优势", before=21.6)
body_para(doc, [("1.", True), ("熟练使用langchain、dify、coze等框架。", False)])
body_para(doc, [("2. 框架应用： 熟悉 Spring Boot、Spring MVC、MyBatis-Plus，能够独立构建 RESTful API。", False)])
body_para(doc, [("3.数据库/缓存： 熟悉 MySQL 索引优化与事务机制；了解 Redis 常用数据结构及其在缓存场景下的应用。", False)])
body_para(doc, [("4.中间件/工具：  熟悉 RabbitMQ 消息队列实现系统解耦与削峰填谷；熟练使用 Docker 容器化部署、"
                 "Git 版本管理及 Linux 常用操作命令", False)])
body_para(doc, [("5. 综合能力： 具备前后端协同开发能力（熟悉 Vue.js），能够快速定位问题并给出解决方案，"
                 "具有良好的代码规范和文档编写习惯。", False)])
body_para(doc, [("6.", False)])

# ---- 项目经历
section_band(doc, "项目经历")
entry_title(doc, "基于 LangChain + LangGraph 的智能影院售票与 AI Agent 助手系统", "2025.12-2026.01")
body_para(doc, [("技术栈：", True), (" Spring Boot, Spring Security, Redis, JWT, AOP, MyBatis-Plus", False)])
body_para(doc, [("权限架构设计： 基于 Spring Security 搭建企业级权限控制中台，设计并实现了多租户模式下的 "
                 "RBAC（基于角色的访问控制）模型，支持细粒度的组织架构权限管理。", False)])
body_para(doc, [("性能优化： 针对权限校验高频访问的痛点，引入 Redis 缓存用户权限信息，通过减少数据库 IO 交互，"
                 "将权限校验接口的响应耗时从 80ms 降低至 10ms 以内，显著减轻了数据库压力。", False)])
body_para(doc, [("动态数据隔离： 为实现不同部门间的数据权限隔离，通过 AOP + MyBatis-Plus 拦截器 实现了动态 SQL "
                 "注入机制，在执行查询前自动过滤部门维度数据，确保了企业级数据的安全性与隔离性。", False)])
body_para(doc, [("灵活校验", True), ("：定制 ", False), ("Spring Security 鉴权逻辑", True),
                ("，实现动态鉴权与访问控制；利用 ", False), ("AOP + MyBatis-Plus", True),
                (" 拦截机制，实现基于部门维度的数据权限过滤，确保业务数据安全性。", False)])
body_para(doc, [("认证体系构建： 基于 JWT 实现无状态认证，并结合 Redis 构建 Token 有效期管理与强制失效机制，"
                 "在保证系统高可用性的同时，解决了分布式环境下的会话同步问题。", False)])

entry_title(doc, "LiliShop", "2026.03-2026.04", before=6.0)
body_para(doc, [("技术栈：", True), (" Spring Boot, Redis, Lua, RabbitMQ, Sentinel, Nginx", False)])
body_para(doc, [("高并发库存处理： 针对秒杀场景下的超卖与数据库压力问题，设计并实现了 “Redis 预减库存 + Lua 脚本” "
                 "方案，将库存扣减逻辑原子化，确保了在高并发环境下库存扣减的绝对准确性,Qps由3000提升至80000。", False)])
body_para(doc, [("异步削峰填谷： 引入 RabbitMQ 构建异步下单链路，将订单创建、支付确认等耗时操作异步化，"
                 "有效解决了瞬时流量冲击导致的系统崩溃，显著提升了系统的吞吐量与响应速度。", False)])
body_para(doc, [("数据一致性保障： 采用 “缓存预热 + 数据库二次校验” 的双重验证机制，通过在异步消费端进行最终库存核验，"
                 "解决了缓存与数据库在极端并发情况下的数据一致性问题。", False)])
body_para(doc, [("系统稳定性增强： 利用 Sentinel 实现细粒度流量控制策略（基于用户 ID + 接口），有效拦截恶意刷单请求；"
                 "配置 Nginx 负载均衡实现动静分离，降低了后端服务器的并发压力。", False)])

# ---- 教育经历
section_band(doc, "教育经历")
edu = doc.add_paragraph()
leading(edu, 20)
edu.paragraph_format.tab_stops.add_tab_stop(Pt(135.8), WD_TAB_ALIGNMENT.LEFT)
edu.paragraph_format.tab_stops.add_tab_stop(Pt(172.7), WD_TAB_ALIGNMENT.LEFT)
edu.paragraph_format.tab_stops.add_tab_stop(Pt(CONTENT_W_PT), WD_TAB_ALIGNMENT.RIGHT)
style_run(edu.add_run("南昌应用技术师范学院"), 12, BODY, bold=True)
style_run(edu.add_run("\t本科"), 9.75, BODY, bold=True)
style_run(edu.add_run("\t软件工程"), 9.75, BODY, bold=True)
style_run(edu.add_run("\t2022-2027"), 9.75, MUTED)

doc.save(OUT)
print("wrote", OUT, os.path.getsize(OUT), "bytes")
